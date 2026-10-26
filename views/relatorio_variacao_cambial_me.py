"""
views/relatorio_variacao_cambial_me.py
=======================================
Relatório de variação cambial para títulos ME em aberto.
Compara a PTAX de faturamento de cada NF com a PTAX BCB de uma data escolhida.
"""
from __future__ import annotations

import datetime
import io

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from components.buttons import botao_exportar_excel
from components.cards import card_html, render_cards_row, render_tabela_html
from utils.ptax import buscar_historico_ptax, buscar_ptax_data
from views.contas_receber_externo import (
    _bloco_cartao, _documento_cartoes, _fmt4, _CORES_PTAX, _sparkline_svg,
)

_SIMBOLOS = {"USD": "$", "EUR": "€", "GBP": "£"}
_TITULOS_CARD = {"USD": "USD/BRL", "EUR": "EUR/BRL", "GBP": "GBP/BRL"}


def _ultimo_dia_util() -> datetime.date:
    d = datetime.date.today() - datetime.timedelta(days=1)
    while d.weekday() >= 5:
        d -= datetime.timedelta(days=1)
    return d


def _parse_nfs(texto: str) -> list[str]:
    if not texto or not texto.strip():
        return []
    return [v.strip() for parte in texto.replace(",", "\n").splitlines() if (v := parte.strip())]


def _fmt_moeda(valor) -> str:
    try:
        v = float(valor)
        return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "R$ 0,00"


def _fmt_num(valor, decimais=2) -> str:
    try:
        return f"{{:,.{decimais}f}}".format(float(valor)).replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "—"


def _chip_variacao(var: float) -> str:
    if var >= 0:
        return '<span style="display:inline-flex;align-items:center;gap:4px;background:rgba(34,197,94,.12);color:#22c55e;border-radius:5px;padding:2px 9px;font-size:11px;font-weight:600">▲ Positiva</span>'
    return '<span style="display:inline-flex;align-items:center;gap:4px;background:rgba(226,75,74,.12);color:#E24B4A;border-radius:5px;padding:2px 9px;font-size:11px;font-weight:600">▼ Negativa</span>'


def _html_var_r(var: float) -> str:
    v = _fmt_moeda(abs(var))
    if var >= 0:
        return f'<span style="color:#22c55e;font-weight:600">▲ {v}</span>'
    return f'<span style="color:#E24B4A;font-weight:600">▼ {v}</span>'


def _html_var_pct(pct: float) -> str:
    v = _fmt_num(abs(pct))
    if pct >= 0:
        return f'<span style="color:#22c55e;font-weight:600">▲ {v}%</span>'
    return f'<span style="color:#E24B4A;font-weight:600">▼ {v}%</span>'


def _card_frente_ptax(moeda: str, ptax_info: dict | None, hist: list) -> str:
    cor = _CORES_PTAX[moeda]
    titulo = _TITULOS_CARD[moeda]
    if not ptax_info:
        return card_html(titulo, "—", subtitulo="sem dados", cor_borda=cor,
                         cor_label="#ffffff", tamanho_valor="20px")
    compra = ptax_info.get("compra", 0)
    subtitulo = f'<span style="color:#6b7280;font-size:11px">compra: R$ {_fmt4(compra)}</span>'
    if hist:
        spark = _sparkline_svg(hist, cor)
        vendas = [h["venda"] for h in hist]
        subtitulo += (
            f'<div style="margin-top:8px;border-top:1px solid #2a2a3e;padding-top:8px">'
            f'{spark}'
            f'<div style="display:flex;justify-content:space-between;margin-top:3px">'
            f'<span style="font-size:10px;color:#555e72">máx/mín 30d</span>'
            f'<span style="font-size:10px;color:#4caf7a">↑ {_fmt4(max(vendas))}</span>'
            f'<span style="font-size:10px;color:#E24B4A">↓ {_fmt4(min(vendas))}</span>'
            f'<a href="https://www.bcb.gov.br/estabilidadefinanceira/historicocotacoes" target="_blank" '
            f'style="font-size:10px;color:{cor};text-decoration:none">BCB ↗</a>'
            f'</div></div>'
        )
    return card_html(titulo, f"R$ {_fmt4(ptax_info['venda'])}", subtitulo=subtitulo,
                     cor_borda=cor, cor_label="#ffffff", tamanho_valor="20px")


def _exportar_excel(df: pd.DataFrame, data_ref: datetime.date, ptax_bcb: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Variação Cambial ME"

    cab_fill   = PatternFill("solid", fgColor="1F3864")
    cab_font   = Font(bold=True, color="FFFFFF", size=10)
    pos_fill   = PatternFill("solid", fgColor="D4EDDA")
    neg_fill   = PatternFill("solid", fgColor="F8D7DA")
    fill_par   = PatternFill("solid", fgColor="FFFFFF")   # linhas pares
    fill_impar = PatternFill("solid", fgColor="EEF4FB")   # linhas ímpares

    data_str = data_ref.strftime('%d/%m/%Y')
    col_ptax_venda = f"PTAX {data_str}"

    colunas  = ["NF", "Parcela", "Entrada", "Container", "Cliente", "País", "Moeda",
                "Saldo em Aberto (Moeda)", "PTAX Fat.",
                "Saldo em Aberto (R$)", col_ptax_venda,
                "Valor Atualizado (R$)", "Variação (R$)",
                "Variação %", "Variação Cambial"]
    larguras = [10, 10, 12, 18, 35, 14, 8, 22, 12, 20, 14, 20, 16, 12, 18]
    num_cols = len(colunas)
    last_col_letter = get_column_letter(num_cols)

    ptax_partes = []
    for moeda, simbolo in (("USD", "$"), ("EUR", "€"), ("GBP", "£")):
        info = (ptax_bcb or {}).get(moeda)
        if info and info.get("venda"):
            ptax_partes.append(f"{simbolo} {_fmt_num(info['venda'], 4)}")
    ptax_str = "  |  ".join(ptax_partes)
    titulo = f"Variação Cambial — Mercado Externo  |  PTAX BCB: {data_str}  |  {ptax_str}" if ptax_str else f"Variação Cambial — Mercado Externo  |  PTAX BCB: {data_str}"

    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=12, color="1F3864")
    ws.row_dimensions[1].height = 22

    for col_idx, (nome, larg) in enumerate(zip(colunas, larguras), 1):
        cel = ws.cell(row=2, column=col_idx, value=nome)
        cel.fill = cab_fill
        cel.font = cab_font
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = larg
    ws.row_dimensions[2].height = 26
    ws.auto_filter.ref = f"A2:{last_col_letter}2"

    fmt2 = '#,##0.00'
    fmt4 = '#,##0.0000'
    fmtp = '#,##0.00"%"'

    # Formato contábil: símbolo à esquerda, número à direita (Excel preenche o espaço com *)
    _fmt_moeda_simb = {
        "USD": '_("$"* #,##0.00_);_("$"* (#,##0.00);_("$"* "-"??_);_(@_)',
        "EUR": '_("€"* #,##0.00_);_("€"* (#,##0.00);_("€"* "-"??_);_(@_)',
        "GBP": '_("£"* #,##0.00_);_("£"* (#,##0.00);_("£"* "-"??_);_(@_)',
    }

    for row_idx, (_, row) in enumerate(df.iterrows(), 3):
        fill = fill_impar if row_idx % 2 != 0 else fill_par
        moeda_cod = str(row.get("MOEDA", ""))
        var_r = float(row.get("VARIACAO_R", 0) or 0)
        var_pct = float(row.get("VARIACAO_PCT", 0) or 0)
        fmt_saldo_moeda = _fmt_moeda_simb.get(moeda_cod, fmt2)

        entrada = pd.to_datetime(row.get("ENTRADA"), errors="coerce")
        container = row.get("CONTAINER")
        pais = row.get("PAIS")
        vals = [
            row.get("NF", ""),
            row.get("PARCELA", ""),
            entrada.to_pydatetime() if pd.notna(entrada) else None,
            container if pd.notna(container) and str(container).strip() else None,
            row.get("CLIENTE", ""),
            pais if pd.notna(pais) and str(pais).strip() else None,
            moeda_cod,
            row.get("SALDO_MOEDA"),
            row.get("COTACAO"),
            row.get("SALDO_EM_ABERTO"),
            row.get("PTAX_BCB"),
            row.get("VL_R_BCB"),
            abs(var_r),
            abs(var_pct),
            "▲ Positiva" if var_r >= 0 else "▼ Negativa",
        ]
        fmts  = [None, None, 'dd/mm/yyyy', None, None, None, None, fmt_saldo_moeda, fmt4, fmt2, fmt4, fmt2, fmt2, fmtp, None]
        aligns = ["left", "center", "center", "left", "left", "left", "center"] + ["right"] * 7 + ["center"]

        for col_idx, (val, fmt, aln) in enumerate(zip(vals, fmts, aligns), 1):
            cel = ws.cell(row=row_idx, column=col_idx, value=val)
            if fmt:
                cel.number_format = fmt
            cel.alignment = Alignment(horizontal=aln)
            cel.fill = fill
            if col_idx in (13, 14, 15):
                cel.font = Font(color="00B050" if var_r >= 0 else "FF0000", bold=True)

    last = ws.max_row + 1
    ws.cell(row=last, column=1, value="TOTAL").font = Font(bold=True)
    for col_idx, campo in [(8, "SALDO_MOEDA"), (10, "SALDO_EM_ABERTO"), (12, "VL_R_BCB")]:
        if campo in df.columns:
            total = df[campo].sum()
            if pd.notna(total):
                cel = ws.cell(row=last, column=col_idx, value=total)
                cel.number_format = fmt2
                cel.font = Font(bold=True)
                cel.alignment = Alignment(horizontal="right")
    var_total = df["VARIACAO_R"].sum()
    cel_var = ws.cell(row=last, column=11, value=abs(var_total))
    cel_var.number_format = fmt2
    cel_var.font = Font(bold=True)
    cel_var.alignment = Alignment(horizontal="right")
    cel_var.fill = pos_fill if var_total >= 0 else neg_fill

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(df: pd.DataFrame):
    if df.empty:
        st.warning("Nenhum dado disponível.")
        return

    df_me = df[
        (df["MERCADO"] == "ME") &
        df["COTACAO"].notna() &
        (df["COTACAO"] > 1)
    ].copy()

    # ── Sidebar ──────────────────────────────────────────────────────────────
    with st.sidebar:
        st.caption("Variação Cambial ME")
        st.markdown(
            "<p style='font-size:17px;font-weight:600;margin:0 0 4px'>📅 Data da PTAX BCB</p>",
            unsafe_allow_html=True,
        )
        data_ptax = st.date_input(
            "📅 Data da PTAX BCB",
            value=_ultimo_dia_util(),
            format="DD/MM/YYYY",
            key="vcm_data_ptax",
            label_visibility="collapsed",
        )
        nfs_texto = st.text_area(
            "Nº das NFs (uma por linha ou vírgula)",
            placeholder="Ex:\n59136\n59138, 59139\n\nVazio = todas em aberto",
            height=130,
            key="vcm_nfs",
        )
        gerar = st.button("Gerar Relatório", width="stretch", key="vcm_btn_gerar")

    if gerar:
        st.session_state["_vcm_gerou"] = True

    # ── Título + botão Atualizar ──────────────────────────────────────────────
    _col_tit, _col_at = st.columns([8, 1.4])
    with _col_tit:
        st.title("Variação Cambial — Mercado Externo")
        st.markdown(
            f'<p style="font-size:13px;color:#1CAE85;font-weight:600;margin:-8px 0 8px">'
            f'📅 PTAX BCB · {data_ptax.strftime("%d/%m/%Y")}</p>',
            unsafe_allow_html=True,
        )
    with _col_at:
        st.markdown("<div style='margin-top:1.05rem'></div>", unsafe_allow_html=True)
        if st.button("🔄 Atualizar", key="vcm_atualizar", use_container_width=True):
            st.cache_data.clear()
            st.rerun()

    _spin = st.empty()
    with _spin.container():
        with st.spinner("Consultando PTAX BCB..."):
            ptax_bcb = buscar_ptax_data(data_ptax)
    _spin.empty()

    data_ref = ptax_bcb.get("data_ref", data_ptax)
    if data_ref != data_ptax:
        st.info(f"PTAX indisponível em {data_ptax.strftime('%d/%m/%Y')} — usando {data_ref.strftime('%d/%m/%Y')}.")

    # ── Cards PTAX com sparkline ──────────────────────────────────────────────
    moedas_disp = [m for m in ("USD", "EUR", "GBP") if ptax_bcb.get(m)]
    hist_cache = {m: buscar_historico_ptax(m, 30) for m in moedas_disp}

    blocos_cot = []
    for i, m in enumerate(moedas_disp):
        frente = _card_frente_ptax(m, ptax_bcb.get(m), hist_cache[m])
        blocos_cot.append(_bloco_cartao(i, m, frente, hist_cache[m]))

    st.iframe(_documento_cartoes(blocos_cot), height=190)

    # ── Aguarda Gerar para exibir tabela ─────────────────────────────────────
    if not st.session_state.get("_vcm_gerou"):
        st.info("Informe as NFs e clique em **Gerar Relatório** para ver a tabela de variação cambial.")
        return

    # ── Filtra NFs ────────────────────────────────────────────────────────────
    nfs = _parse_nfs(st.session_state.get("vcm_nfs", ""))
    df_rel = df_me[df_me["NF"].astype(str).isin(nfs)].copy() if nfs else df_me.copy()

    if df_rel.empty:
        st.warning("Nenhum título ME em aberto encontrado para as NFs informadas.")
        return

    # ── Calcula variação ──────────────────────────────────────────────────────
    df_rel["PTAX_BCB"] = df_rel["MOEDA"].map(
        lambda m: ptax_bcb[m]["venda"] if ptax_bcb.get(m) else None
    )
    df_rel["VL_R_BCB"]    = df_rel["SALDO_MOEDA"] * df_rel["PTAX_BCB"]
    df_rel["VARIACAO_R"]  = df_rel["VL_R_BCB"] - df_rel["SALDO_EM_ABERTO"]
    df_rel["VARIACAO_PCT"] = (
        df_rel["VARIACAO_R"] / df_rel["SALDO_EM_ABERTO"].abs() * 100
    ).where(df_rel["SALDO_EM_ABERTO"] != 0)

    st.divider()

    # ── Tabela ────────────────────────────────────────────────────────────────
    col_ptax_venda = f"PTAX {data_ref.strftime('%d/%m/%Y')}"

    cols_base = ["NF", "PARCELA", "ENTRADA", "CONTAINER", "CLIENTE", "PAIS", "MOEDA",
                 "SALDO_MOEDA", "COTACAO", "SALDO_EM_ABERTO",
                 "PTAX_BCB", "VL_R_BCB", "VARIACAO_R", "VARIACAO_PCT"]
    for c in cols_base:
        if c not in df_rel.columns:
            df_rel[c] = None

    df_exib = df_rel[cols_base].copy()

    df_exib["MOEDA"]           = df_rel["MOEDA"].astype(str)
    df_exib["ENTRADA"]         = pd.to_datetime(df_rel["ENTRADA"], errors="coerce").dt.strftime("%d/%m/%Y").fillna("—")
    df_exib["CONTAINER"]       = df_rel["CONTAINER"].apply(lambda v: str(v) if pd.notna(v) and str(v).strip() else "—")
    df_exib["PAIS"]            = df_rel["PAIS"].apply(lambda v: str(v) if pd.notna(v) and str(v).strip() else "—")
    df_exib["SALDO_MOEDA"]     = df_rel.apply(
        lambda r: (
            f'<span style="display:flex;justify-content:space-between;gap:8px">'
            f'<span>{_SIMBOLOS.get(str(r["MOEDA"]), str(r["MOEDA"]))}</span>'
            f'<span>{_fmt_num(r["SALDO_MOEDA"])}</span></span>'
        ) if pd.notna(r["SALDO_MOEDA"]) else "—", axis=1
    )
    df_exib["COTACAO"]         = df_rel["COTACAO"].apply(lambda v: _fmt_num(v, 4) if pd.notna(v) else "—")
    df_exib["SALDO_EM_ABERTO"] = df_rel["SALDO_EM_ABERTO"].apply(
        lambda v: (
            f'<span style="display:flex;justify-content:space-between;gap:8px">'
            f'<span>R$</span><span>{_fmt_num(v)}</span></span>'
        ) if pd.notna(v) else "—"
    )
    df_exib["PTAX_BCB"]        = df_rel["PTAX_BCB"].apply(lambda v: _fmt_num(v, 4) if pd.notna(v) else "—")
    df_exib["VL_R_BCB"]        = df_rel["VL_R_BCB"].apply(
        lambda v: (
            f'<span style="display:flex;justify-content:space-between;gap:8px">'
            f'<span>R$</span><span>{_fmt_num(v)}</span></span>'
        ) if pd.notna(v) else "—"
    )
    df_exib["VARIACAO_R"]      = df_rel["VARIACAO_R"].apply(
        lambda v: _html_var_r(v) if pd.notna(v) else "—"
    )
    df_exib["VARIACAO_PCT"]    = df_rel["VARIACAO_PCT"].apply(
        lambda v: _html_var_pct(v) if pd.notna(v) else "—"
    )
    df_exib["VARIACAO_CAMBIAL"] = df_rel["VARIACAO_R"].apply(
        lambda v: _chip_variacao(v) if pd.notna(v) else "—"
    )

    df_exib = df_exib.rename(columns={
        "PARCELA":          "Parcela",
        "ENTRADA":          "Entrada",
        "CONTAINER":        "Container",
        "CLIENTE":          "Cliente",
        "PAIS":             "País",
        "SALDO_MOEDA":      "Saldo em Aberto (Moeda)",
        "COTACAO":          "PTAX Fat.",
        "SALDO_EM_ABERTO":  "Saldo em Aberto (R$)",
        "PTAX_BCB":         col_ptax_venda,
        "VL_R_BCB":         "Valor Atualizado (R$)",
        "VARIACAO_R":       "Variação (R$)",
        "VARIACAO_PCT":     "Variação %",
        "VARIACAO_CAMBIAL": "Variação Cambial",
    })

    render_tabela_html(
        df_exib,
        colunas_centralizar=["MOEDA", "Parcela", "Entrada", "PTAX Fat.", col_ptax_venda, "Variação Cambial"],
    )

    # ── Exportar Excel ────────────────────────────────────────────────────────
    excel_bytes = _exportar_excel(df_rel, data_ref, ptax_bcb)
    st.markdown("<div style='margin-top:16px'></div>", unsafe_allow_html=True)
    col_xl, _ = st.columns([1, 5])
    with col_xl:
        botao_exportar_excel(
            excel_bytes,
            f"Variacao Cambial ME {data_ref.strftime('%d-%m-%Y')}.xlsx",
            key="vcm_export",
            use_container_width=True,
        )
